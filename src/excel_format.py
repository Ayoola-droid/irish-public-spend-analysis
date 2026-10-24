from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
EURO = '€#,##0.00'
EURO_COLUMNS = {"euro", "spend", "calculated", "stated_total", "difference", "extra_value",
                "same_quarter_last_year", "previous_quarter", "latest_quarter",
                "change_vs_last_year"}


def style_sheet(ws):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for col in ws.columns:
        header = str(col[0].value)
        width = max(len(str(c.value)) for c in col if c.value is not None)
        if header in EURO_COLUMNS:
            width += 6
            for c in col[1:]:
                c.number_format = EURO
        ws.column_dimensions[col[0].column_letter].width = min(width + 5, 60)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def add_bar_chart(ws, title, value_col, anchor, rows=None, horizontal=False):
    last = rows + 1 if rows else ws.max_row
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.title = title
    chart.title.overlay = False
    chart.legend = None
    chart.varyColors = False
    chart.height, chart.width = (12, 20) if horizontal else (9, 18)
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.y_axis.number_format = '€#,##0,,"m"'
    data = Reference(ws, min_col=value_col, min_row=1, max_row=last)
    labels = Reference(ws, min_col=1, min_row=2, max_row=last)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(labels)
    if horizontal:
        chart.x_axis.axPos = "l"
        chart.y_axis.axPos = "b"
        chart.x_axis.scaling.orientation = "maxMin"
        chart.y_axis.crosses = "max"
    else:
        chart.x_axis.axPos = "b"
    ws.add_chart(chart, anchor)


def format_workbook(path):
    wb = load_workbook(path)
    for ws in wb.worksheets:
        style_sheet(ws)

    ws = wb["Summary"]
    for measure, value in ws.iter_rows(min_row=2):
        if "(EUR)" in str(measure.value):
            value.number_format = EURO
    ws.auto_filter.ref = None

    add_bar_chart(wb["By Quarter"], "Spend by quarter (EUR)", 3, "E2")
    add_bar_chart(wb["Top Suppliers"], "Top 10 suppliers by spend (EUR)", 3, "H2",
                  rows=10, horizontal=True)
    add_bar_chart(wb["By Category"], "Spend by category (EUR)", 3, "F2", horizontal=True)
    wb.save(path)